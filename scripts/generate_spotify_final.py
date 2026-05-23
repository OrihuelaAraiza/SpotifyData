from __future__ import annotations

import csv
import html
import json
import math
import random
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from statistics import median
from zoneinfo import ZoneInfo

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.backends.backend_pdf import PdfPages
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
PROJECT = Path(__file__).resolve().parents[1]
DATASETS = [
    {
        "person_id": "juan_pablo",
        "name": "Juan Pablo Orihuela Araiza",
        "folder": ROOT / "Spotify Extended Streaming History",
    },
    {
        "person_id": "aranza",
        "name": "Aranza Romo Lima",
        "folder": ROOT / "Spotify Extended Streaming History 2",
    },
]
TZ_NAME = "America/Mexico_City"
TZ = ZoneInfo(TZ_NAME)
SESSION_GAP_MINUTES = 30
ANOMALOUS_GROUP_THRESHOLD = 3
WEEKDAYS = ["Lun", "Mar", "Mie", "Jue", "Vie", "Sab", "Dom"]
PALETTE = {
    "green": "#1DB954",
    "ink": "#191414",
    "muted": "#63706A",
    "blue": "#247BA0",
    "red": "#D1495B",
    "gold": "#EDAE49",
    "teal": "#2A9D8F",
    "purple": "#6D597A",
    "paper": "#F7F8F5",
}


def out(*parts: str) -> Path:
    return PROJECT.joinpath(*parts)


def ensure_dirs() -> None:
    for path in [
        out("data", "processed"),
        out("data", "processed", "combinado"),
        out("notebooks"),
        out("reports", "figures", "comparativo"),
        out("reports", "personas"),
        out("scripts"),
        out("docs"),
    ]:
        path.mkdir(parents=True, exist_ok=True)
    for cfg in DATASETS:
        pid = cfg["person_id"]
        out("data", "processed", pid).mkdir(parents=True, exist_ok=True)
        out("reports", "figures", pid).mkdir(parents=True, exist_ok=True)


def safe_text(value: object, default: str = "") -> str:
    if value is None:
        return default
    text = str(value).strip()
    return text if text else default


def parse_ts(value: str) -> datetime:
    return datetime.fromisoformat(value.replace("Z", "+00:00"))


def time_block(hour: int) -> str:
    if 5 <= hour <= 11:
        return "manana"
    if 12 <= hour <= 17:
        return "tarde"
    if 18 <= hour <= 23:
        return "noche"
    return "madrugada"


def platform_bucket(value: str | None) -> str:
    if not value:
        return "Other"
    text = value.lower()
    if "iphone" in text or text == "ios" or ("ios" in text and "ipad" not in text):
        return "iPhone / iOS"
    if "ipad" in text:
        return "iPad / iOS"
    if "android" in text:
        return "Android"
    if "web" in text:
        return "Web Player"
    if "windows" in text:
        return "Windows"
    if "mac" in text or "os x" in text:
        return "macOS"
    if "linux" in text:
        return "Linux"
    return "Other"


def row_signature(row: dict) -> tuple:
    return (
        row.get("ts"),
        row.get("platform"),
        row.get("ms_played"),
        row.get("conn_country"),
        row.get("master_metadata_track_name"),
        row.get("master_metadata_album_artist_name"),
        row.get("master_metadata_album_album_name"),
        row.get("spotify_track_uri"),
        row.get("episode_name"),
        row.get("episode_show_name"),
        row.get("spotify_episode_uri"),
        row.get("audiobook_title"),
        row.get("audiobook_uri"),
        row.get("audiobook_chapter_uri"),
        row.get("audiobook_chapter_title"),
        row.get("reason_start"),
        row.get("reason_end"),
        row.get("shuffle"),
        row.get("skipped"),
        row.get("offline"),
        row.get("offline_timestamp"),
        row.get("incognito_mode"),
    )


def load_and_clean(cfg: dict) -> tuple[list[dict], dict]:
    folder = cfg["folder"]
    audio_files = sorted(folder.glob("Streaming_History_Audio_*.json"))
    video_files = sorted(folder.glob("Streaming_History_Video_*.json"))
    grouped: dict[tuple[str | None, str | None], dict[tuple, dict]] = defaultdict(dict)
    rows_by_file = {}
    raw_rows = 0
    exact_duplicates = 0
    nulls = Counter()
    fields = set()

    for path in audio_files:
        data = json.loads(path.read_text(encoding="utf-8"))
        rows_by_file[path.name] = len(data)
        for row in data:
            raw_rows += 1
            fields.update(row.keys())
            for key, value in row.items():
                if value in (None, ""):
                    nulls[key] += 1
            row = dict(row)
            row["source_file"] = path.name
            key = (row.get("ts"), row.get("platform"))
            signature = row_signature(row)
            if signature in grouped[key]:
                exact_duplicates += 1
                continue
            grouped[key][signature] = row

    cleaned = []
    collapsed_groups = 0
    collapsed_rows = 0
    for signatures in grouped.values():
        values = list(signatures.values())
        if len(values) > ANOMALOUS_GROUP_THRESHOLD:
            collapsed_groups += 1
            collapsed_rows += len(values) - 1
            cleaned.append(max(values, key=lambda r: int(r.get("ms_played") or 0)))
        else:
            cleaned.extend(values)

    cleaned.sort(key=lambda r: r.get("ts") or "")
    video_rows = 0
    for path in video_files:
        try:
            video_rows += len(json.loads(path.read_text(encoding="utf-8")))
        except json.JSONDecodeError:
            pass

    stats = {
        "person_id": cfg["person_id"],
        "name": cfg["name"],
        "source_folder": str(folder),
        "audio_files": len(audio_files),
        "video_files": len(video_files),
        "video_rows_available": video_rows,
        "rows_by_file": rows_by_file,
        "raw_rows": raw_rows,
        "exact_duplicates_removed": exact_duplicates,
        "collapsed_groups": collapsed_groups,
        "collapsed_rows_removed": collapsed_rows,
        "cleaned_rows": len(cleaned),
        "field_names": sorted(f for f in fields if f != "ip_addr"),
        "field_nulls": {k: v for k, v in sorted(nulls.items()) if k != "ip_addr"},
    }
    return cleaned, stats


def enrich_rows(rows: list[dict], cfg: dict) -> list[dict]:
    enriched = []
    previous_dt = None
    session_id = 0
    for row in rows:
        if not row.get("ts"):
            continue
        dt_utc = parse_ts(row["ts"])
        dt_local = dt_utc.astimezone(TZ)
        if previous_dt is None or (dt_utc - previous_dt).total_seconds() / 60 > SESSION_GAP_MINUTES:
            session_id += 1
        previous_dt = dt_utc
        ms = int(row.get("ms_played") or 0)
        track = safe_text(row.get("master_metadata_track_name"))
        artist = safe_text(row.get("master_metadata_album_artist_name"))
        album = safe_text(row.get("master_metadata_album_album_name"))
        episode = safe_text(row.get("episode_name"))
        content_type = "track" if row.get("spotify_track_uri") else "episode" if row.get("spotify_episode_uri") else "other"
        enriched.append(
            {
                "person_id": cfg["person_id"],
                "person_name": cfg["name"],
                "event_id": len(enriched) + 1,
                "session_id": session_id,
                "ts_utc": dt_utc.isoformat(),
                "datetime_local": dt_local.strftime("%Y-%m-%d %H:%M:%S"),
                "date_local": dt_local.date().isoformat(),
                "year": dt_local.year,
                "year_month": dt_local.strftime("%Y-%m"),
                "iso_week": f"{dt_local.isocalendar()[0]}-W{dt_local.isocalendar()[1]:02d}",
                "month": dt_local.month,
                "weekday": WEEKDAYS[dt_local.weekday()],
                "weekday_num": dt_local.weekday(),
                "hour": dt_local.hour,
                "is_weekend": dt_local.weekday() >= 5,
                "time_block": time_block(dt_local.hour),
                "platform_raw": safe_text(row.get("platform"), "Unknown"),
                "device_group": platform_bucket(row.get("platform")),
                "conn_country": safe_text(row.get("conn_country"), "Unknown"),
                "content_type": content_type,
                "artist": artist if artist else "Unknown Artist",
                "track": track if track else episode if episode else "Unknown Content",
                "album": album if album else "Unknown Album",
                "spotify_track_uri": safe_text(row.get("spotify_track_uri")),
                "ms_played": ms,
                "minutes_played": round(ms / 60000, 4),
                "played_30s_or_more": ms >= 30000,
                "skipped": bool(row.get("skipped")),
                "completed_or_ended_naturally": row.get("reason_end") in {"trackdone", "endplay"},
                "shuffle": bool(row.get("shuffle")),
                "offline": bool(row.get("offline")),
                "incognito_mode": bool(row.get("incognito_mode")),
                "reason_start": safe_text(row.get("reason_start"), "Unknown"),
                "reason_end": safe_text(row.get("reason_end"), "Unknown"),
                "source_file": safe_text(row.get("source_file"), "Unknown"),
            }
        )
    return enriched


def pct(part: float, total: float) -> float:
    return 0.0 if not total else part / total * 100.0


def build_sessions(rows: list[dict]) -> list[dict]:
    grouped = defaultdict(list)
    for row in rows:
        if row["ms_played"] > 0:
            grouped[row["session_id"]].append(row)
    sessions = []
    for session_id, items in sorted(grouped.items()):
        total_ms = sum(i["ms_played"] for i in items)
        artists = {i["artist"] for i in items if i["content_type"] == "track"}
        tracks = {(i["artist"], i["track"]) for i in items if i["content_type"] == "track"}
        skipped = sum(1 for i in items if i["skipped"])
        device = Counter(i["device_group"] for i in items).most_common(1)[0][0]
        first, last = items[0], items[-1]
        sessions.append(
            {
                "person_id": first["person_id"],
                "person_name": first["person_name"],
                "session_id": session_id,
                "start_local": first["datetime_local"],
                "end_local": last["datetime_local"],
                "date_local": first["date_local"],
                "weekday": first["weekday"],
                "hour_start": first["hour"],
                "time_block_start": first["time_block"],
                "device_group": device,
                "tracks_played": len(items),
                "unique_artists": len(artists),
                "unique_tracks": len(tracks),
                "minutes_listened": round(total_ms / 60000, 4),
                "skip_rate_pct": round(pct(skipped, len(items)), 2),
            }
        )
    return sessions


def aggregate(rows: list[dict], sessions: list[dict], cleaning: dict) -> dict:
    total_ms = sum(r["ms_played"] for r in rows)
    track_rows = [r for r in rows if r["content_type"] == "track"]
    active_days = {r["date_local"] for r in rows if r["ms_played"] > 0}
    artists = {r["artist"] for r in track_rows}
    tracks = {(r["artist"], r["track"]) for r in track_rows}
    artist_ms, artist_plays = Counter(), Counter()
    track_ms, track_plays = Counter(), Counter()
    month_ms, month_plays, year_ms = Counter(), Counter(), Counter()
    hour_ms, hour_rows, hour_skips = Counter(), Counter(), Counter()
    weekday_ms, weekday_hour_ms = Counter(), Counter()
    platform_ms, country_rows = Counter(), Counter()
    completion, reason_start, reason_end = Counter(), Counter(), Counter()
    month_track_rows = defaultdict(list)
    seen_tracks = set()
    discovery = defaultdict(lambda: {"plays": 0, "new_tracks": 0})

    for r in rows:
        ms = r["ms_played"]
        month_ms[r["year_month"]] += ms
        month_plays[r["year_month"]] += 1
        year_ms[r["year"]] += ms
        hour_ms[r["hour"]] += ms
        hour_rows[r["hour"]] += 1
        weekday_ms[r["weekday"]] += ms
        weekday_hour_ms[(r["weekday_num"], r["hour"])] += ms
        platform_ms[r["device_group"]] += ms
        country_rows[r["conn_country"]] += 1
        reason_start[r["reason_start"]] += 1
        reason_end[r["reason_end"]] += 1
        if r["skipped"]:
            completion["saltadas"] += 1
            hour_skips[r["hour"]] += 1
        elif r["completed_or_ended_naturally"]:
            completion["completas"] += 1
        else:
            completion["parciales"] += 1
        if r["content_type"] == "track":
            key = (r["artist"], r["track"])
            artist_ms[r["artist"]] += ms
            artist_plays[r["artist"]] += 1
            track_ms[key] += ms
            track_plays[key] += 1
            month_track_rows[r["year_month"]].append(r)
            discovery[r["year_month"]]["plays"] += 1
            if key not in seen_tracks:
                discovery[r["year_month"]]["new_tracks"] += 1
                seen_tracks.add(key)

    monthly_features = []
    for month in sorted(month_ms):
        rows_m = month_track_rows[month]
        total_month_ms = month_ms[month]
        artist_month_ms = Counter()
        track_month_plays = Counter()
        skipped = 0
        night_ms = 0
        for r in rows_m:
            artist_month_ms[r["artist"]] += r["ms_played"]
            track_month_plays[(r["artist"], r["track"])] += 1
            skipped += 1 if r["skipped"] else 0
            night_ms += r["ms_played"] if r["time_block"] in {"noche", "madrugada"} else 0
        monthly_features.append(
            {
                "person_id": cleaning["person_id"],
                "person_name": cleaning["name"],
                "year_month": month,
                "hours": round(total_month_ms / 3600000, 4),
                "plays": month_plays[month],
                "unique_artists": len({r["artist"] for r in rows_m}),
                "unique_tracks": len({(r["artist"], r["track"]) for r in rows_m}),
                "skip_rate_pct": round(pct(skipped, len(rows_m)), 4) if rows_m else 0.0,
                "night_ratio_pct": round(pct(night_ms, total_month_ms), 4) if total_month_ms else 0.0,
                "top_artist_share_pct": round(pct(artist_month_ms.most_common(1)[0][1], total_month_ms), 4) if artist_month_ms else 0.0,
                "repeat_track_ratio_pct": round(pct(sum(1 for v in track_month_plays.values() if v > 1), len(track_month_plays)), 4) if track_month_plays else 0.0,
                "new_tracks": discovery[month]["new_tracks"],
                "discovery_rate_pct": round(pct(discovery[month]["new_tracks"], discovery[month]["plays"]), 4) if discovery[month]["plays"] else 0.0,
            }
        )

    top_artists = [
        {"artist": a, "hours": round(ms / 3600000, 2), "plays": artist_plays[a], "share_pct": round(pct(ms, total_ms), 2)}
        for a, ms in artist_ms.most_common(20)
    ]
    top_tracks = [
        {"artist": a, "track": t, "hours": round(ms / 3600000, 2), "plays": track_plays[(a, t)], "share_pct": round(pct(ms, total_ms), 3)}
        for (a, t), ms in track_ms.most_common(20)
    ]
    session_minutes = [s["minutes_listened"] for s in sessions]
    max_month, max_month_ms = max(month_ms.items(), key=lambda kv: kv[1])
    max_year, max_year_ms = max(year_ms.items(), key=lambda kv: kv[1])
    max_hour = max(range(24), key=lambda h: hour_ms[h])
    max_weekday = max(WEEKDAYS, key=lambda d: weekday_ms[d])
    top10_ms = sum(v for _, v in artist_ms.most_common(10))
    top20_ms = sum(v for _, v in artist_ms.most_common(20))
    min_dt = min(parse_ts(r["ts_utc"]) for r in rows)
    max_dt = max(parse_ts(r["ts_utc"]) for r in rows)

    summary = {
        "metadata": {
            **cleaning,
            "timezone": TZ_NAME,
            "date_min_local": min(r["datetime_local"] for r in rows),
            "date_max_local": max(r["datetime_local"] for r in rows),
            "date_min_utc": min_dt.isoformat(),
            "date_max_utc": max_dt.isoformat(),
        },
        "totals": {
            "total_hours": round(total_ms / 3600000, 2),
            "total_days_listening": round(total_ms / 86400000, 2),
            "active_days": len(active_days),
            "avg_hours_per_active_day": round((total_ms / 3600000) / len(active_days), 2),
            "track_rows": len(track_rows),
            "episode_rows": sum(1 for r in rows if r["content_type"] == "episode"),
            "unique_artists": len(artists),
            "unique_tracks": len(tracks),
            "skip_rate_pct": round(pct(sum(1 for r in rows if r["skipped"]), len(rows)), 2),
            "weekend_share_pct": round(pct(sum(r["ms_played"] for r in rows if r["is_weekend"]), total_ms), 2),
        },
        "peaks": {
            "peak_hour": max_hour,
            "peak_hour_hours": round(hour_ms[max_hour] / 3600000, 2),
            "peak_weekday": max_weekday,
            "peak_weekday_hours": round(weekday_ms[max_weekday] / 3600000, 2),
            "peak_month": max_month,
            "peak_month_hours": round(max_month_ms / 3600000, 2),
            "peak_year": max_year,
            "peak_year_hours": round(max_year_ms / 3600000, 2),
        },
        "concentration": {
            "top_10_artist_share_pct": round(pct(top10_ms, total_ms), 2),
            "top_20_artist_share_pct": round(pct(top20_ms, total_ms), 2),
            "top_artist_share_pct": top_artists[0]["share_pct"] if top_artists else 0.0,
        },
        "sessions": {
            "count": len(sessions),
            "avg_listening_minutes": round(sum(session_minutes) / len(session_minutes), 2),
            "median_listening_minutes": round(float(median(session_minutes)), 2),
            "p90_listening_minutes": round(float(np.percentile(session_minutes, 90)), 2),
            "avg_tracks_per_session": round(sum(s["tracks_played"] for s in sessions) / len(sessions), 2),
        },
        "top_artists": top_artists,
        "top_tracks": top_tracks,
        "platforms": [{"platform": k, "hours": round(v / 3600000, 2), "share_pct": round(pct(v, total_ms), 2)} for k, v in platform_ms.most_common()],
        "countries": [{"country": k, "rows": v, "share_pct": round(pct(v, len(rows)), 2)} for k, v in country_rows.most_common(8)],
        "completion": dict(completion),
        "reason_start": dict(reason_start.most_common(10)),
        "reason_end": dict(reason_end.most_common(10)),
        "series": {
            "monthly_hours": {m: round(month_ms[m] / 3600000, 2) for m in sorted(month_ms)},
            "hourly_hours": {str(h): round(hour_ms[h] / 3600000, 2) for h in range(24)},
            "hourly_skip_rate": {str(h): round(pct(hour_skips[h], hour_rows[h]), 2) if hour_rows[h] else 0.0 for h in range(24)},
            "weekday_hours": {d: round(weekday_ms[d] / 3600000, 2) for d in WEEKDAYS},
            "weekday_hour_hours": {f"{WEEKDAYS[d]}_{h:02d}": round(weekday_hour_ms[(d, h)] / 3600000, 2) for d in range(7) for h in range(24)},
            "monthly_features": monthly_features,
        },
    }
    summary["ml"] = run_ml(monthly_features, rows)
    return summary


def standardize(matrix: np.ndarray) -> np.ndarray:
    std = matrix.std(axis=0)
    std[std == 0] = 1
    return (matrix - matrix.mean(axis=0)) / std


def kmeans(matrix: np.ndarray, k: int = 3, iterations: int = 100) -> tuple[np.ndarray, float, float]:
    rng = np.random.default_rng(42)
    centers = matrix[rng.choice(len(matrix), size=min(k, len(matrix)), replace=False)]
    if len(centers) < k:
        return np.zeros(len(matrix), dtype=int), 0.0, 0.0
    labels = np.zeros(len(matrix), dtype=int)
    for _ in range(iterations):
        distances = ((matrix[:, None, :] - centers[None, :, :]) ** 2).sum(axis=2)
        new_labels = distances.argmin(axis=1)
        new_centers = np.array([matrix[new_labels == i].mean(axis=0) if np.any(new_labels == i) else centers[i] for i in range(k)])
        if np.array_equal(new_labels, labels):
            centers = new_centers
            break
        labels, centers = new_labels, new_centers
    inertia = float(sum(((matrix[i] - centers[labels[i]]) ** 2).sum() for i in range(len(matrix))))
    return labels, inertia, silhouette(matrix, labels)


def silhouette(matrix: np.ndarray, labels: np.ndarray) -> float:
    unique = sorted(set(labels.tolist()))
    if len(unique) < 2 or len(unique) >= len(matrix):
        return 0.0
    scores = []
    for idx, point in enumerate(matrix):
        same = matrix[labels == labels[idx]]
        other = [matrix[labels == label] for label in unique if label != labels[idx]]
        a = np.mean(np.linalg.norm(same - point, axis=1)) if len(same) > 1 else 0.0
        b = min(np.mean(np.linalg.norm(cluster - point, axis=1)) for cluster in other if len(cluster))
        scores.append((b - a) / max(a, b) if max(a, b) else 0.0)
    return float(np.mean(scores))


def run_ml(monthly_features: list[dict], rows: list[dict]) -> dict:
    features = ["hours", "unique_artists", "unique_tracks", "skip_rate_pct", "night_ratio_pct", "top_artist_share_pct", "repeat_track_ratio_pct", "discovery_rate_pct"]
    matrix = np.array([[float(r[f]) for f in features] for r in monthly_features], dtype=float)
    labels, inertia, sil = kmeans(standardize(matrix), k=3)
    counts = Counter(labels.tolist())
    profiles = []
    for label in sorted(counts):
        cluster = [monthly_features[i] for i, value in enumerate(labels) if value == label]
        profile = {"cluster": int(label), "months": len(cluster)}
        for f in features:
            profile[f"avg_{f}"] = round(sum(float(r[f]) for r in cluster) / len(cluster), 3)
        profiles.append(profile)
    return {
        "clustering": {
            "model": "K-Means mensual implementado con NumPy",
            "k": 3,
            "features": features,
            "inertia": round(inertia, 4),
            "silhouette": round(sil, 4),
            "cluster_counts": {str(k): v for k, v in sorted(counts.items())},
            "profiles": profiles,
        },
        "skip_classifier": train_skip_classifier(rows),
    }


def train_skip_classifier(rows: list[dict]) -> dict:
    sample = [r for r in rows if r["content_type"] == "track"]
    top_artists = {a for a, _ in Counter(r["artist"] for r in sample).most_common(25)}
    platforms = sorted({p for p, _ in Counter(r["device_group"] for r in sample).most_common(6)})
    reasons = sorted({r for r, _ in Counter(row["reason_start"] for row in sample).most_common(8)})
    rng = random.Random(42)
    rng.shuffle(sample)
    sample = sample[: min(len(sample), 60000)]
    feature_names = ["bias", "hour_sin", "hour_cos", "is_weekend", "shuffle", "offline", "artist_top25"] + [f"platform_{p}" for p in platforms] + [f"reason_start_{r}" for r in reasons]
    x = np.array([encode_skip_features(r, top_artists, platforms, reasons) for r in sample], dtype=float)
    y = np.array([1.0 if r["skipped"] else 0.0 for r in sample], dtype=float)
    x[:, 1:] = standardize(x[:, 1:])
    split = int(len(x) * 0.8)
    x_train, x_test = x[:split], x[split:]
    y_train, y_test = y[:split], y[split:]
    weights = np.zeros(x_train.shape[1], dtype=float)
    for _ in range(650):
        pred = 1 / (1 + np.exp(-np.clip(x_train @ weights, -40, 40)))
        weights -= 0.08 * ((x_train.T @ (pred - y_train)) / len(y_train))
    probs = 1 / (1 + np.exp(-np.clip(x_test @ weights, -40, 40)))
    y_pred = (probs >= 0.5).astype(float)
    tp = int(((y_pred == 1) & (y_test == 1)).sum())
    tn = int(((y_pred == 0) & (y_test == 0)).sum())
    fp = int(((y_pred == 1) & (y_test == 0)).sum())
    fn = int(((y_pred == 0) & (y_test == 1)).sum())
    accuracy = (tp + tn) / len(y_test)
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    coefficients = sorted(
        [{"feature": name, "coefficient": round(float(value), 4)} for name, value in zip(feature_names, weights)],
        key=lambda item: abs(item["coefficient"]),
        reverse=True,
    )[:10]
    return {
        "model": "Regresion logistica binaria implementada con NumPy",
        "target": "skipped",
        "sample_rows": len(sample),
        "train_rows": len(y_train),
        "test_rows": len(y_test),
        "accuracy": round(accuracy, 4),
        "precision": round(precision, 4),
        "recall": round(recall, 4),
        "f1": round(f1, 4),
        "confusion_matrix": {"tp": tp, "tn": tn, "fp": fp, "fn": fn},
        "top_coefficients_abs": coefficients,
    }


def encode_skip_features(row: dict, top_artists: set[str], platforms: list[str], reasons: list[str]) -> list[float]:
    angle = 2 * math.pi * row["hour"] / 24
    values = [
        1.0,
        math.sin(angle),
        math.cos(angle),
        1.0 if row["is_weekend"] else 0.0,
        1.0 if row["shuffle"] else 0.0,
        1.0 if row["offline"] else 0.0,
        1.0 if row["artist"] in top_artists else 0.0,
    ]
    values.extend(1.0 if row["device_group"] == p else 0.0 for p in platforms)
    values.extend(1.0 if row["reason_start"] == r else 0.0 for r in reasons)
    return values


def write_csv(path: Path, rows: list[dict], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fields})


STREAM_FIELDS = [
    "person_id", "person_name", "event_id", "session_id", "ts_utc", "datetime_local", "date_local", "year", "year_month", "iso_week", "month", "weekday", "weekday_num", "hour", "is_weekend", "time_block", "platform_raw", "device_group", "conn_country", "content_type", "artist", "track", "album", "spotify_track_uri", "ms_played", "minutes_played", "played_30s_or_more", "skipped", "completed_or_ended_naturally", "shuffle", "offline", "incognito_mode", "reason_start", "reason_end", "source_file"
]
SESSION_FIELDS = ["person_id", "person_name", "session_id", "start_local", "end_local", "date_local", "weekday", "hour_start", "time_block_start", "device_group", "tracks_played", "unique_artists", "unique_tracks", "minutes_listened", "skip_rate_pct"]
MONTHLY_FIELDS = ["person_id", "person_name", "year_month", "hours", "plays", "unique_artists", "unique_tracks", "skip_rate_pct", "night_ratio_pct", "top_artist_share_pct", "repeat_track_ratio_pct", "new_tracks", "discovery_rate_pct"]


def write_processed(cfg: dict, rows: list[dict], sessions: list[dict], summary: dict) -> None:
    base = out("data", "processed", cfg["person_id"])
    write_csv(base / "spotify_streams_limpio.csv", rows, STREAM_FIELDS)
    write_csv(base / "spotify_sessions.csv", sessions, SESSION_FIELDS)
    write_csv(base / "spotify_monthly_features.csv", summary["series"]["monthly_features"], MONTHLY_FIELDS)
    (base / "spotify_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")


def write_combined_processed(rows: list[dict], sessions: list[dict], monthly: list[dict]) -> None:
    base = out("data", "processed", "combinado")
    write_csv(base / "spotify_streams_limpio_combinado.csv", rows, STREAM_FIELDS)
    write_csv(base / "spotify_sessions_combinado.csv", sessions, SESSION_FIELDS)
    write_csv(base / "spotify_monthly_features_combinado.csv", monthly, MONTHLY_FIELDS)


def horizontal_bar(title: str, labels: list[str], values: list[float], output: Path, color: str) -> None:
    height = max(5, 0.34 * len(labels))
    fig, ax = plt.subplots(figsize=(11, height))
    ax.barh(labels[::-1], values[::-1], color=color)
    ax.set_title(title)
    ax.set_xlabel("Horas")
    ax.grid(axis="x", alpha=0.25)
    fig.tight_layout()
    fig.savefig(output, dpi=170)
    plt.close(fig)


def make_person_figures(summary: dict) -> None:
    pid = summary["metadata"]["person_id"]
    figdir = out("reports", "figures", pid)
    monthly = summary["series"]["monthly_hours"]
    labels = list(monthly)
    values = [monthly[k] for k in labels]
    rolling = [sum(values[max(0, i - 2): i + 1]) / len(values[max(0, i - 2): i + 1]) for i in range(len(values))]
    fig, ax = plt.subplots(figsize=(14, 5))
    ax.plot(labels, values, color=PALETTE["green"], linewidth=1.4, label="Horas")
    ax.plot(labels, rolling, color=PALETTE["gold"], linewidth=2.0, label="Media movil 3m")
    ax.set_title("Horas escuchadas por mes")
    ax.set_ylabel("Horas")
    ax.set_xticks(range(0, len(labels), max(1, len(labels) // 12)))
    ax.set_xticklabels(labels[:: max(1, len(labels) // 12)], rotation=45, ha="right")
    ax.grid(alpha=0.25)
    ax.legend()
    fig.tight_layout()
    fig.savefig(figdir / "01_horas_por_mes.png", dpi=170)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(10, 4.6))
    ax.bar(range(24), [summary["series"]["hourly_hours"][str(h)] for h in range(24)], color=PALETTE["teal"])
    ax.set_title("Distribucion de escucha por hora local")
    ax.set_xlabel("Hora")
    ax.set_ylabel("Horas")
    ax.set_xticks(range(24))
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    fig.savefig(figdir / "02_escucha_por_hora.png", dpi=170)
    plt.close(fig)

    heatmap = np.zeros((7, 24))
    for d, day in enumerate(WEEKDAYS):
        for h in range(24):
            heatmap[d, h] = summary["series"]["weekday_hour_hours"][f"{day}_{h:02d}"]
    fig, ax = plt.subplots(figsize=(12, 4.8))
    img = ax.imshow(heatmap, aspect="auto", cmap="YlGnBu")
    ax.set_title("Mapa de calor: dia x hora")
    ax.set_yticks(range(7))
    ax.set_yticklabels(WEEKDAYS)
    ax.set_xticks(range(24))
    fig.colorbar(img, ax=ax, label="Horas")
    fig.tight_layout()
    fig.savefig(figdir / "03_heatmap_semana_hora.png", dpi=170)
    plt.close(fig)

    horizontal_bar("Top 20 artistas por horas", [x["artist"] for x in summary["top_artists"]], [x["hours"] for x in summary["top_artists"]], figdir / "04_top20_artistas_horas.png", PALETTE["red"])
    horizontal_bar("Top 20 canciones por horas", [f"{x['track']} - {x['artist']}" for x in summary["top_tracks"]], [x["hours"] for x in summary["top_tracks"]], figdir / "05_top20_canciones_horas.png", PALETTE["blue"])

    mf = summary["series"]["monthly_features"]
    fig, ax = plt.subplots(figsize=(14, 5))
    ax.plot([r["year_month"] for r in mf], [r["unique_artists"] for r in mf], color=PALETTE["purple"], label="Artistas unicos")
    ax.plot([r["year_month"] for r in mf], [r["unique_tracks"] for r in mf], color=PALETTE["green"], label="Canciones unicas")
    ax.set_title("Diversidad musical mensual")
    ax.set_ylabel("Conteo unico")
    ax.set_xticks(range(0, len(mf), max(1, len(mf) // 12)))
    ax.set_xticklabels([r["year_month"] for r in mf][:: max(1, len(mf) // 12)], rotation=45, ha="right")
    ax.grid(alpha=0.25)
    ax.legend()
    fig.tight_layout()
    fig.savefig(figdir / "06_diversidad_mensual.png", dpi=170)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(7, 4.6))
    completion = summary["completion"]
    ax.bar(completion.keys(), completion.values(), color=[PALETTE["green"], PALETTE["gold"], PALETTE["red"]])
    ax.set_title("Completacion vs skip")
    ax.set_ylabel("Reproducciones")
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    fig.savefig(figdir / "07_completion_vs_skip.png", dpi=170)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(10, 4.6))
    ax.bar(range(24), [summary["series"]["hourly_skip_rate"][str(h)] for h in range(24)], color=PALETTE["red"])
    ax.set_title("Tasa de skip por hora")
    ax.set_xlabel("Hora")
    ax.set_ylabel("Skip rate (%)")
    ax.set_xticks(range(24))
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    fig.savefig(figdir / "08_skip_rate_por_hora.png", dpi=170)
    plt.close(fig)

    session_csv = out("data", "processed", pid, "spotify_sessions.csv")
    minutes = [float(r["minutes_listened"]) for r in csv.DictReader(session_csv.open(encoding="utf-8"))]
    upper = np.percentile(minutes, 95)
    fig, ax = plt.subplots(figsize=(10, 4.6))
    ax.hist([m for m in minutes if m <= upper], bins=30, color=PALETTE["purple"])
    ax.axvline(float(np.median(minutes)), color=PALETTE["red"], linestyle="--", label="Mediana")
    ax.set_title("Distribucion de duracion de sesiones")
    ax.set_xlabel("Minutos escuchados")
    ax.set_ylabel("Sesiones")
    ax.grid(axis="y", alpha=0.25)
    ax.legend()
    fig.tight_layout()
    fig.savefig(figdir / "09_sesiones_histograma.png", dpi=170)
    plt.close(fig)

    horizontal_bar("Horas por dispositivo", [x["platform"] for x in summary["platforms"]], [x["hours"] for x in summary["platforms"]], figdir / "10_dispositivos.png", PALETTE["blue"])


def make_comparison_figures(summaries: list[dict]) -> None:
    figdir = out("reports", "figures", "comparativo")
    all_months = sorted(set().union(*(set(s["series"]["monthly_hours"]) for s in summaries)))
    fig, ax = plt.subplots(figsize=(14, 5))
    colors = [PALETTE["green"], PALETTE["blue"]]
    for idx, summary in enumerate(summaries):
        monthly = summary["series"]["monthly_hours"]
        ax.plot(all_months, [monthly.get(m, 0) for m in all_months], label=summary["metadata"]["name"].split()[0], color=colors[idx], linewidth=1.5)
    ax.set_title("Comparativo de horas escuchadas por mes")
    ax.set_ylabel("Horas")
    ax.set_xticks(range(0, len(all_months), max(1, len(all_months) // 12)))
    ax.set_xticklabels(all_months[:: max(1, len(all_months) // 12)], rotation=45, ha="right")
    ax.grid(alpha=0.25)
    ax.legend()
    fig.tight_layout()
    fig.savefig(figdir / "01_comparativo_horas_mes.png", dpi=170)
    plt.close(fig)

    metrics = ["total_hours", "unique_artists", "unique_tracks", "skip_rate_pct"]
    labels = ["Horas", "Artistas unicos", "Canciones unicas", "Skip %"]
    x = np.arange(len(metrics))
    width = 0.35
    fig, ax = plt.subplots(figsize=(10, 5))
    for idx, summary in enumerate(summaries):
        vals = [summary["totals"][m] for m in metrics]
        ax.bar(x + (idx - 0.5) * width, vals, width, label=summary["metadata"]["name"].split()[0], color=colors[idx])
    ax.set_title("Indicadores generales comparados")
    ax.set_xticks(x)
    ax.set_xticklabels(labels)
    ax.grid(axis="y", alpha=0.25)
    ax.legend()
    fig.tight_layout()
    fig.savefig(figdir / "02_comparativo_kpis.png", dpi=170)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(10, 5))
    for idx, summary in enumerate(summaries):
        ax.plot(range(24), [summary["series"]["hourly_hours"][str(h)] for h in range(24)], marker="o", label=summary["metadata"]["name"].split()[0], color=colors[idx])
    ax.set_title("Comparativo de rutina por hora")
    ax.set_xlabel("Hora local")
    ax.set_ylabel("Horas")
    ax.set_xticks(range(24))
    ax.grid(alpha=0.25)
    ax.legend()
    fig.tight_layout()
    fig.savefig(figdir / "03_comparativo_hora.png", dpi=170)
    plt.close(fig)


def build_combined_summary(summaries: list[dict]) -> dict:
    return {
        "project": "Proyecto Final Spotify",
        "timezone": TZ_NAME,
        "people": summaries,
        "comparison": {
            "most_hours": max(summaries, key=lambda s: s["totals"]["total_hours"])["metadata"]["name"],
            "highest_diversity_tracks": max(summaries, key=lambda s: s["totals"]["unique_tracks"])["metadata"]["name"],
            "highest_skip_rate": max(summaries, key=lambda s: s["totals"]["skip_rate_pct"])["metadata"]["name"],
            "total_clean_rows": sum(s["metadata"]["cleaned_rows"] for s in summaries),
            "total_raw_rows": sum(s["metadata"]["raw_rows"] for s in summaries),
        },
    }


def wrap(text: str, width: int = 92) -> str:
    words = text.split()
    lines, current = [], []
    for word in words:
        if sum(len(w) + 1 for w in current) + len(word) > width and current:
            lines.append(" ".join(current))
            current = [word]
        else:
            current.append(word)
    if current:
        lines.append(" ".join(current))
    return "\n".join(lines)


def pdf_text_page(pdf: PdfPages, title: str, bullets: list[str]) -> None:
    fig = plt.figure(figsize=(8.5, 11))
    fig.patch.set_facecolor("white")
    fig.text(0.08, 0.94, title, fontsize=20, fontweight="bold", color=PALETTE["ink"])
    fig.text(0.08, 0.915, "Proyecto Final Spotify - Exploracion y Comprension de los Datos", fontsize=9, color=PALETTE["muted"])
    y = 0.86
    for bullet in bullets:
        text = wrap(bullet, 86)
        fig.text(0.1, y, "- " + text.replace("\n", "\n  "), fontsize=11.5, color=PALETTE["ink"], va="top")
        y -= 0.05 * (text.count("\n") + 1)
    pdf.savefig(fig)
    plt.close(fig)


def pdf_image_page(pdf: PdfPages, title: str, image: Path, note: str) -> None:
    fig = plt.figure(figsize=(8.5, 11))
    fig.patch.set_facecolor("white")
    fig.text(0.08, 0.94, title, fontsize=19, fontweight="bold", color=PALETTE["ink"])
    ax = fig.add_axes([0.08, 0.21, 0.84, 0.62])
    ax.axis("off")
    ax.imshow(Image.open(image))
    fig.text(0.08, 0.12, wrap(note, 95), fontsize=11, color=PALETTE["ink"])
    pdf.savefig(fig)
    plt.close(fig)


def write_pdf(summaries: list[dict], combined: dict) -> None:
    pdf_path = out("reports", "spotify_informe_ejecutivo.pdf")
    juan, aranza = summaries
    with PdfPages(pdf_path) as pdf:
        pdf_text_page(pdf, "Informe ejecutivo Spotify", [
            "Alcance: solo Spotify, con dos historiales extendidos: Juan Pablo Orihuela Araiza y Aranza Romo Lima.",
            f"Base total procesada: {combined['comparison']['total_raw_rows']:,} filas crudas y {combined['comparison']['total_clean_rows']:,} filas limpias.",
            "Se entregan notebooks por fase, CSV limpios, figuras, dashboard HTML, informe ejecutivo PDF y presentacion.",
            "La capa final omite direcciones IP y reporta patrones agregados para reducir exposicion de datos sensibles.",
        ])
        pdf_text_page(pdf, "Metodologia", [
            "Se trabajo con una arquitectura Bronze/Silver/Gold: JSON crudos referenciados, CSV limpios por persona y resumen analitico para reportes.",
            "Los timestamps UTC se convirtieron a America/Mexico_City. Se derivaron dia, hora, mes, semana ISO, bloque horario y fin de semana.",
            "Las sesiones se definieron con gaps mayores a 30 minutos entre reproducciones.",
            "La concentracion musical se mide por artista, ya que el export local de Spotify no incluye genero nativo.",
        ])
        pdf_text_page(pdf, "Limpieza y calidad", [
            f"Juan Pablo: {juan['metadata']['raw_rows']:,} filas crudas, {juan['metadata']['cleaned_rows']:,} limpias, {juan['metadata']['exact_duplicates_removed']:,} duplicados exactos removidos.",
            f"Aranza: {aranza['metadata']['raw_rows']:,} filas crudas, {aranza['metadata']['cleaned_rows']:,} limpias, {aranza['metadata']['exact_duplicates_removed']:,} duplicados exactos removidos.",
            "Las rafagas anomalas con mismo timestamp y plataforma se colapsaron conservando el evento de mayor duracion.",
            "Los campos sensibles como IP no se exportan en los CSV finales.",
        ])
        pdf_image_page(pdf, "Comparativo mensual", out("reports", "figures", "comparativo", "01_comparativo_horas_mes.png"), "La serie mensual permite comparar intensidad y cambios de etapa entre ambas personas.")
        pdf_image_page(pdf, "Indicadores comparados", out("reports", "figures", "comparativo", "02_comparativo_kpis.png"), f"Mayor tiempo total: {combined['comparison']['most_hours']}. Mayor diversidad por canciones: {combined['comparison']['highest_diversity_tracks']}.")
        pdf_image_page(pdf, "Rutina por hora", out("reports", "figures", "comparativo", "03_comparativo_hora.png"), "La distribucion por hora muestra diferencias en habitos diarios y momentos pico.")
        for summary in summaries:
            name = summary["metadata"]["name"].split()[0]
            pid = summary["metadata"]["person_id"]
            pdf_image_page(pdf, f"{name}: patron temporal", out("reports", "figures", pid, "01_horas_por_mes.png"), f"Mes pico: {summary['peaks']['peak_month']} con {summary['peaks']['peak_month_hours']:.1f} horas.")
            pdf_image_page(pdf, f"{name}: top artistas", out("reports", "figures", pid, "04_top20_artistas_horas.png"), f"Artista principal: {summary['top_artists'][0]['artist']} con {summary['top_artists'][0]['hours']:.1f} horas y {summary['top_artists'][0]['plays']:,} reproducciones.")
            pdf_image_page(pdf, f"{name}: diversidad mensual", out("reports", "figures", pid, "06_diversidad_mensual.png"), f"Artistas unicos: {summary['totals']['unique_artists']:,}; canciones unicas: {summary['totals']['unique_tracks']:,}.")
        pdf_text_page(pdf, "Resultados de Machine Learning", [
            f"Juan Pablo, K-Means mensual: silhouette {juan['ml']['clustering']['silhouette']:.3f}; clasificador de skip F1 {juan['ml']['skip_classifier']['f1']:.3f}.",
            f"Aranza, K-Means mensual: silhouette {aranza['ml']['clustering']['silhouette']:.3f}; clasificador de skip F1 {aranza['ml']['skip_classifier']['f1']:.3f}.",
            "Modelo 1: K-Means segmenta meses por intensidad, diversidad, skip, nocturnidad, concentracion y descubrimiento.",
            "Modelo 2: regresion logistica estima probabilidad de skip usando hora, fin de semana, shuffle, offline, artista frecuente, dispositivo y razon de inicio.",
        ])
        pdf_text_page(pdf, "Privacidad y conclusiones", [
            "Spotify permite inferir rutinas, dispositivos, paises de conexion, preferencias persistentes y momentos de cambio personal.",
            "Los resultados muestran que el historial musical no es solo entretenimiento: funciona como huella digital temporal y conductual.",
            "Para mejorar el proyecto se podria enriquecer con generos o audio features desde una API externa, documentando el costo de privacidad y dependencia tecnica.",
            "La entrega final cumple limpieza, EDA, ML, reflexion, visualizaciones y organizacion de codigo solicitadas en el PDF.",
        ])


def write_markdown_reports(summaries: list[dict], combined: dict) -> None:
    lines = [
        "# Checklist de requerimientos Spotify",
        "",
        "| Requisito PDF | Evidencia | Estado |",
        "|---|---|---|",
        "| Fase 1: estructura, JSON principales, descripcion y reflexion | `notebooks/01_exploracion_spotify.ipynb`, `reports/spotify_informe_ejecutivo.pdf` | Completo |",
        "| Fase 2: carga, calidad, nulos/duplicados, CSV limpio y transformaciones | `notebooks/02_limpieza_consolidacion.ipynb`, `data/processed/*`, `reports/informe_transformaciones.md` | Completo |",
        "| Fase 3: EDA con visualizaciones, insights y comparacion por periodos | `notebooks/03_eda_spotify.ipynb`, `reports/informe_eda.md`, `reports/figures` | Completo |",
        "| Patrones temporales: hora, dia, mensual, heatmap y picos | Figuras 01, 02, 03 por persona y comparativo | Completo |",
        "| Totales de artistas y canciones unicas | `reports/summary_combined.json`, informe PDF | Completo |",
        "| Top 20 artistas y canciones | Figuras 04 y 05 por persona | Completo |",
        "| Diversidad musical por mes | Figura 06 y `spotify_monthly_features.csv` | Completo |",
        "| Completion vs skip y sesiones | Figuras 07, 08, 09 y CSV de sesiones | Completo |",
        "| Concentracion por generos o artistas | Concentracion por artistas, porque el export no trae genero nativo | Completo |",
        "| ML: al menos 2 modelos, metricas e interpretacion | `notebooks/04_ml_spotify.ipynb`, informe PDF | Completo |",
        "| Entrega final: PDF 10-15 paginas y presentacion 5-10 minutos | `reports/spotify_informe_ejecutivo.pdf`, `reports/spotify_presentacion.pptx` | Completo |",
        "| Privacidad | CSV sin IP, resultados agregados | Completo |",
        "",
        f"Conteos validados: {combined['comparison']['total_raw_rows']:,} filas crudas y {combined['comparison']['total_clean_rows']:,} filas limpias entre las dos personas.",
    ]
    out("reports", "checklist_requerimientos_spotify.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

    transform = ["# Informe de transformaciones Spotify", ""]
    for summary in summaries:
        c = summary["metadata"]
        transform += [
            f"## {c['name']}",
            f"- Carpeta fuente: `{c['source_folder']}`.",
            f"- JSON de audio leidos: {c['audio_files']}. JSON de video detectados: {c['video_files']} con {c['video_rows_available']:,} registros disponibles.",
            f"- Filas crudas de audio: {c['raw_rows']:,}.",
            f"- Duplicados exactos removidos: {c['exact_duplicates_removed']:,}.",
            f"- Grupos anomalos timestamp+platform colapsados: {c['collapsed_groups']:,}.",
            f"- Filas removidas por colapso conservador: {c['collapsed_rows_removed']:,}.",
            f"- Filas limpias finales: {c['cleaned_rows']:,}.",
            "",
        ]
    transform += [
        "## Transformaciones comunes",
        "- Conversion de `ts` UTC a `datetime_local` en `America/Mexico_City`.",
        "- Derivacion de fecha, mes, semana ISO, dia de semana, hora, bloque horario y fin de semana.",
        "- Normalizacion de plataforma en categorias de dispositivo.",
        "- Clasificacion de contenido en track, episode u other.",
        "- Construccion de sesiones con separacion mayor a 30 minutos.",
        "- Exportacion sin direcciones IP.",
    ]
    out("reports", "informe_transformaciones.md").write_text("\n".join(transform) + "\n", encoding="utf-8")

    eda = ["# Informe EDA Spotify", ""]
    for summary in summaries:
        t, p, c, s = summary["totals"], summary["peaks"], summary["concentration"], summary["sessions"]
        eda += [
            f"## {summary['metadata']['name']}",
            f"- Ventana analizada: {summary['metadata']['date_min_local']} a {summary['metadata']['date_max_local']}.",
            f"- Escucha total: {t['total_hours']:,.2f} horas en {t['active_days']:,} dias activos.",
            f"- Artistas unicos: {t['unique_artists']:,}. Canciones unicas: {t['unique_tracks']:,}.",
            f"- Artista mas escuchado: {summary['top_artists'][0]['artist']} ({summary['top_artists'][0]['plays']:,} reproducciones).",
            f"- Cancion mas escuchada: {summary['top_tracks'][0]['track']} - {summary['top_tracks'][0]['artist']} ({summary['top_tracks'][0]['plays']:,} reproducciones).",
            f"- Hora pico: {p['peak_hour']:02d}:00. Dia pico: {p['peak_weekday']}. Mes pico: {p['peak_month']}.",
            f"- Skip rate global: {t['skip_rate_pct']:.2f}%. Top 10 artistas concentran {c['top_10_artist_share_pct']:.2f}% del tiempo.",
            f"- Sesiones detectadas: {s['count']:,}; mediana {s['median_listening_minutes']:.1f} minutos; P90 {s['p90_listening_minutes']:.1f} minutos.",
            "",
        ]
    eda += [
        "## Comparacion",
        f"- Mayor tiempo total: {combined['comparison']['most_hours']}.",
        f"- Mayor diversidad de canciones: {combined['comparison']['highest_diversity_tracks']}.",
        f"- Mayor tasa de skip: {combined['comparison']['highest_skip_rate']}.",
        "",
        "## Reflexion",
        "El historial de escucha permite identificar rutinas y preferencias con mucho detalle. Aunque los datos parecen recreativos, combinan horario, dispositivo, pais y conducta de consumo; por eso se reportan agregados y se omiten identificadores sensibles.",
    ]
    out("reports", "informe_eda.md").write_text("\n".join(eda) + "\n", encoding="utf-8")

    anexo = [
        "# Anexo visual Spotify",
        "",
        "La fase multimedia es opcional y el historial extendido de Spotify no incluye imagenes, portadas ni fotos locales. En lugar de inventar multimedia, se entrega un dashboard visual con patrones temporales, top contenido, diversidad, completion/skip, sesiones y comparativos.",
        "",
        "Archivo principal: `reports/dashboard_spotify.html`.",
    ]
    out("reports", "anexo_visual_spotify.md").write_text("\n".join(anexo) + "\n", encoding="utf-8")


def write_dashboard(summaries: list[dict], combined: dict) -> None:
    cards = []
    for s in summaries:
        pid = s["metadata"]["person_id"]
        cards.append(f"""
        <section>
          <h2>{html.escape(s['metadata']['name'])}</h2>
          <div class="kpis">
            <div><b>{s['totals']['total_hours']:,.0f}</b><span>horas</span></div>
            <div><b>{s['totals']['unique_artists']:,}</b><span>artistas</span></div>
            <div><b>{s['totals']['unique_tracks']:,}</b><span>canciones</span></div>
            <div><b>{s['totals']['skip_rate_pct']:.1f}%</b><span>skip</span></div>
          </div>
          <div class="grid">
            <img src="figures/{pid}/01_horas_por_mes.png" alt="Horas por mes">
            <img src="figures/{pid}/03_heatmap_semana_hora.png" alt="Heatmap semana hora">
            <img src="figures/{pid}/04_top20_artistas_horas.png" alt="Top artistas">
            <img src="figures/{pid}/07_completion_vs_skip.png" alt="Completion vs skip">
          </div>
        </section>
        """)
    html_doc = f"""<!doctype html>
<html lang="es">
<head>
<meta charset="utf-8">
<title>Dashboard Spotify - Proyecto Final</title>
<style>
body {{ margin:0; font-family: Arial, sans-serif; background:#f7f8f5; color:#191414; }}
header {{ padding:32px 44px; background:#191414; color:white; }}
h1 {{ margin:0 0 8px; font-size:34px; }}
h2 {{ margin:34px 0 16px; }}
main {{ max-width:1180px; margin:0 auto; padding:28px; }}
.kpis {{ display:grid; grid-template-columns: repeat(4, 1fr); gap:12px; margin:16px 0 22px; }}
.kpis div {{ background:white; border:1px solid #d8ddd8; border-radius:8px; padding:14px; }}
.kpis b {{ display:block; font-size:26px; color:#1DB954; }}
.kpis span {{ color:#63706A; }}
.grid {{ display:grid; grid-template-columns: repeat(2, minmax(0,1fr)); gap:18px; }}
img {{ width:100%; background:white; border:1px solid #d8ddd8; border-radius:8px; }}
section {{ margin-bottom:40px; }}
@media (max-width:800px) {{ .grid,.kpis {{ grid-template-columns:1fr; }} header {{ padding:24px; }} main {{ padding:18px; }} }}
</style>
</head>
<body>
<header>
<h1>Proyecto Final Spotify</h1>
<p>Dos historiales extendidos, limpieza, EDA, Machine Learning y reflexion de privacidad.</p>
</header>
<main>
<section>
<h2>Comparativo</h2>
<div class="kpis">
<div><b>{combined['comparison']['total_clean_rows']:,}</b><span>filas limpias</span></div>
<div><b>{html.escape(combined['comparison']['most_hours'].split()[0])}</b><span>mayor tiempo</span></div>
<div><b>{html.escape(combined['comparison']['highest_diversity_tracks'].split()[0])}</b><span>mayor diversidad</span></div>
<div><b>{html.escape(combined['comparison']['highest_skip_rate'].split()[0])}</b><span>mayor skip</span></div>
</div>
<div class="grid">
<img src="figures/comparativo/01_comparativo_horas_mes.png" alt="Comparativo mensual">
<img src="figures/comparativo/02_comparativo_kpis.png" alt="Comparativo KPIs">
<img src="figures/comparativo/03_comparativo_hora.png" alt="Comparativo por hora">
</div>
</section>
{''.join(cards)}
</main>
</body>
</html>
"""
    out("reports", "dashboard_spotify.html").write_text(html_doc, encoding="utf-8")


def add_textbox(slide, text: str, left, top, width, height, size=22, bold=False, color="191414"):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return box


def add_bullets(slide, bullets: list[str], left, top, width, height, size=18):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    for idx, bullet in enumerate(bullets):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor.from_string("191414")
    return box


def write_pptx(summaries: list[dict], combined: dict) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def slide(title: str, bullets: list[str], image: Path | None = None):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = RGBColor.from_string("F7F8F5")
        add_textbox(sl, title, Inches(0.55), Inches(0.35), Inches(12), Inches(0.55), size=28, bold=True)
        if image:
            sl.shapes.add_picture(str(image), Inches(0.8), Inches(1.25), width=Inches(11.8), height=Inches(4.45))
            add_bullets(sl, bullets, Inches(0.8), Inches(6.0), Inches(11.6), Inches(0.8), size=14)
        else:
            add_bullets(sl, bullets, Inches(0.9), Inches(1.45), Inches(11.4), Inches(4.8), size=21)

    slide("Spotify: dos huellas musicales comparadas", [
        f"{combined['comparison']['total_clean_rows']:,} eventos limpios",
        "Juan Pablo Orihuela Araiza y Aranza Romo Lima",
        "Limpieza, EDA, Machine Learning, dashboard y reflexion de privacidad",
    ])
    slide("Base procesada con trazabilidad", [
        f"Juan Pablo: {summaries[0]['metadata']['audio_files']} JSON de audio, {summaries[0]['metadata']['cleaned_rows']:,} filas limpias",
        f"Aranza: {summaries[1]['metadata']['audio_files']} JSON de audio, {summaries[1]['metadata']['cleaned_rows']:,} filas limpias",
        "CSV final sin direcciones IP",
    ])
    slide("Intensidad mensual comparada", ["Cambios de etapa y meses pico visibles en el tiempo."], out("reports", "figures", "comparativo", "01_comparativo_horas_mes.png"))
    slide("Indicadores generales", [f"Mayor tiempo: {combined['comparison']['most_hours']}", f"Mayor diversidad: {combined['comparison']['highest_diversity_tracks']}"], out("reports", "figures", "comparativo", "02_comparativo_kpis.png"))
    for s in summaries:
        pid = s["metadata"]["person_id"]
        first = s["metadata"]["name"].split()[0]
        slide(f"{first}: favoritos principales", [f"Top artista: {s['top_artists'][0]['artist']}", f"Top 10 artistas = {s['concentration']['top_10_artist_share_pct']:.1f}% del tiempo"], out("reports", "figures", pid, "04_top20_artistas_horas.png"))
        slide(f"{first}: diversidad y exploracion", [f"{s['totals']['unique_tracks']:,} canciones unicas", f"Mes pico: {s['peaks']['peak_month']}"], out("reports", "figures", pid, "06_diversidad_mensual.png"))
    slide("Machine Learning", [
        f"Juan Pablo: K-Means silhouette {summaries[0]['ml']['clustering']['silhouette']:.3f}; skip F1 {summaries[0]['ml']['skip_classifier']['f1']:.3f}",
        f"Aranza: K-Means silhouette {summaries[1]['ml']['clustering']['silhouette']:.3f}; skip F1 {summaries[1]['ml']['skip_classifier']['f1']:.3f}",
        "Los modelos resumen comportamiento; no automatizan decisiones personales.",
    ])
    slide("Privacidad y aprendizaje", [
        "Spotify revela rutinas, dispositivos, paises y preferencias persistentes.",
        "Se reportan agregados y se omiten campos sensibles.",
        "La huella musical tambien es una huella conductual.",
    ])
    prs.save(out("reports", "spotify_presentacion.pptx"))


def write_notebooks() -> None:
    notebooks = {
        "01_exploracion_spotify.ipynb": [
            ("markdown", "# Fase 1: Exploracion y comprension\n\nIdentifica estructura, JSON principales, contenido y riesgos de privacidad de ambos historiales de Spotify."),
            ("code", "from pathlib import Path\nimport json\nbase = Path('..')\nfor folder in ['Spotify Extended Streaming History','Spotify Extended Streaming History 2']:\n    files = sorted((base / folder).glob('Streaming_History_Audio_*.json'))\n    video = sorted((base / folder).glob('Streaming_History_Video_*.json'))\n    rows = sum(len(json.load(open(path, encoding='utf-8'))) for path in files)\n    print(folder, 'audio_json=', len(files), 'video_json=', len(video), 'filas_audio=', rows)"),
            ("markdown", "Reflexion: el historial permite inferir horarios, dispositivos, paises de conexion, gustos persistentes y cambios de rutina. Por eso el proyecto exporta resultados agregados y no direcciones IP."),
        ],
        "02_limpieza_consolidacion.ipynb": [
            ("markdown", "# Fase 2: Limpieza y consolidacion\n\nCarga, calidad, duplicados, normalizacion temporal y exportacion de CSV limpios."),
            ("code", "import sys\nfrom pathlib import Path\nsys.path.append(str((Path('..') / 'scripts').resolve()))\nimport generate_spotify_final as pipeline\npipeline.main()"),
            ("markdown", "Transformaciones: UTC a America/Mexico_City, variables temporales, normalizacion de dispositivo, clasificacion de contenido, sesiones con gaps mayores a 30 minutos, remocion de duplicados y exportacion sin IP."),
        ],
        "03_eda_spotify.ipynb": [
            ("markdown", "# Fase 3: EDA Spotify\n\nPatrones temporales, top contenido, diversidad, completion/skip, sesiones y comparativo."),
            ("code", "import json\nsummary = json.load(open('../reports/summary_combined.json', encoding='utf-8'))\nfor person in summary['people']:\n    print('\\n', person['metadata']['name'])\n    print('Totales:', person['totals'])\n    print('Picos:', person['peaks'])\n    print('Top artista:', person['top_artists'][0])\n    print('Top cancion:', person['top_tracks'][0])"),
            ("code", "from IPython.display import Image, display\nfor path in ['../reports/figures/comparativo/01_comparativo_horas_mes.png','../reports/figures/comparativo/02_comparativo_kpis.png','../reports/figures/juan_pablo/04_top20_artistas_horas.png','../reports/figures/aranza/04_top20_artistas_horas.png']:\n    display(Image(filename=path))"),
        ],
        "04_ml_spotify.ipynb": [
            ("markdown", "# Fase 4: Machine Learning\n\nDos modelos por persona: K-Means mensual y regresion logistica para skip."),
            ("code", "import json\nsummary = json.load(open('../reports/summary_combined.json', encoding='utf-8'))\nfor person in summary['people']:\n    print('\\n', person['metadata']['name'])\n    print('K-Means:', person['ml']['clustering'])\n    print('Skip classifier:', person['ml']['skip_classifier'])"),
            ("markdown", "Interpretacion: K-Means agrupa meses por intensidad, diversidad, skip, nocturnidad, concentracion y descubrimiento. La regresion logistica muestra que saltar canciones tiene senales observables, pero no es completamente determinista."),
        ],
    }
    for filename, cells in notebooks.items():
        data = {
            "cells": [],
            "metadata": {"kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"}, "language_info": {"name": "python", "version": "3"}},
            "nbformat": 4,
            "nbformat_minor": 5,
        }
        for cell_type, source in cells:
            cell = {"cell_type": cell_type, "metadata": {}, "source": [line + "\n" for line in source.splitlines()]}
            if cell_type == "code":
                cell.update({"execution_count": None, "outputs": []})
            data["cells"].append(cell)
        out("notebooks", filename).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def write_readme(combined: dict) -> None:
    text = f"""# Proyecto Final Spotify

Proyecto generado para dos personas:

- Juan Pablo Orihuela Araiza
- Aranza Romo Lima

## Estructura

- `data/processed/`: CSV limpios por persona y resumen JSON.
- `data/processed/combinado/`: CSV limpios unidos para comparar ambas personas.
- `notebooks/`: notebooks organizados por fase.
- `reports/figures/`: visualizaciones individuales y comparativas.
- `reports/spotify_informe_ejecutivo.pdf`: informe ejecutivo de 10-15 paginas.
- `reports/spotify_presentacion.pptx`: presentacion para 5-10 minutos.
- `reports/dashboard_spotify.html`: dashboard explicativo.
- `scripts/generate_spotify_final.py`: pipeline reproducible.

## Reproducir

```bash
source .venv/bin/activate
python scripts/generate_spotify_final.py
```

El pipeline usa los JSON fuente ubicados en la carpeta padre del proyecto y no exporta direcciones IP.

Filas procesadas: {combined['comparison']['total_raw_rows']:,} crudas; {combined['comparison']['total_clean_rows']:,} limpias.
"""
    out("README.md").write_text(text, encoding="utf-8")
    out("requirements.txt").write_text("numpy\nmatplotlib\npillow\npypdf\npython-pptx\n", encoding="utf-8")


def copy_assignment_pdf() -> None:
    source = ROOT / "Instagram y Spotify.pdf"
    target = out("docs", "Instagram y Spotify.pdf")
    if source.exists():
        target.write_bytes(source.read_bytes())


def main() -> None:
    ensure_dirs()
    summaries = []
    all_rows = []
    all_sessions = []
    all_monthly = []
    for cfg in DATASETS:
        cleaned, cleaning = load_and_clean(cfg)
        rows = enrich_rows(cleaned, cfg)
        sessions = build_sessions(rows)
        summary = aggregate(rows, sessions, cleaning)
        write_processed(cfg, rows, sessions, summary)
        make_person_figures(summary)
        summaries.append(summary)
        all_rows.extend(rows)
        all_sessions.extend(sessions)
        all_monthly.extend(summary["series"]["monthly_features"])
    write_combined_processed(all_rows, all_sessions, all_monthly)
    make_comparison_figures(summaries)
    combined = build_combined_summary(summaries)
    out("reports", "summary_combined.json").write_text(json.dumps(combined, ensure_ascii=False, indent=2), encoding="utf-8")
    write_markdown_reports(summaries, combined)
    write_dashboard(summaries, combined)
    write_pdf(summaries, combined)
    write_pptx(summaries, combined)
    write_notebooks()
    write_readme(combined)
    copy_assignment_pdf()
    print(f"Proyecto generado en: {PROJECT}")
    print(f"Filas limpias totales: {combined['comparison']['total_clean_rows']:,}")
    print(f"Informe: {out('reports', 'spotify_informe_ejecutivo.pdf')}")
    print(f"Presentacion: {out('reports', 'spotify_presentacion.pptx')}")


if __name__ == "__main__":
    main()
